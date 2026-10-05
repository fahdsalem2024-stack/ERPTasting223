"""
Export Helper - PDF, Excel, JSON
"""
import json
from datetime import datetime
from io import BytesIO
from pathlib import Path


def export_to_json(summaries):
    """Export reports to JSON"""
    output = {
        "exported_at": datetime.now().isoformat(),
        "total_reports": len(summaries),
        "reports": summaries,
    }
    return json.dumps(output, ensure_ascii=False, indent=2).encode("utf-8")


def export_to_csv(summaries):
    """Export reports to CSV"""
    import csv
    from io import StringIO

    output = StringIO()
    writer = csv.writer(output)

    # Header
    writer.writerow([
        "Run ID", "Test Name", "Total", "Passed", "Failed",
        "Duration (s)", "Timestamp", "Exit Code"
    ])

    # Rows
    for s in summaries:
        writer.writerow([
            s.get("run_id", ""),
            s.get("test_name", ""),
            s.get("total", 0),
            s.get("passed", 0),
            s.get("failed", 0) + s.get("errors", 0),
            int(s.get("duration_seconds", 0)),
            s.get("timestamp", ""),
            s.get("exit_code", ""),
        ])

    # Add BOM for Excel to read Arabic
    content = "\ufeff" + output.getvalue()
    return content.encode("utf-8")


def export_to_excel(summaries):
    """Export reports to Excel with styling"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Reports"

    # Headers
    headers = [
        "Run ID", "Test Name", "Total", "Passed", "Failed",
        "Duration (s)", "Timestamp", "Exit Code"
    ]

    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border

    # Data rows
    success_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    fail_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

    for row_idx, s in enumerate(summaries, 2):
        failed = s.get("failed", 0) + s.get("errors", 0)
        row_data = [
            s.get("run_id", ""),
            s.get("test_name", ""),
            s.get("total", 0),
            s.get("passed", 0),
            failed,
            int(s.get("duration_seconds", 0)),
            s.get("timestamp", ""),
            s.get("exit_code", ""),
        ]
        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if col == 5:  # Failed column
                cell.fill = fail_fill if failed > 0 else success_fill

    # Column widths
    widths = [20, 30, 10, 10, 10, 15, 20, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    # Summary sheet
    ws2 = wb.create_sheet("Summary")
    total_runs = len(summaries)
    total_passed = sum(s.get("passed", 0) for s in summaries)
    total_failed = sum(s.get("failed", 0) + s.get("errors", 0) for s in summaries)
    success_rate = round((total_passed / (total_passed + total_failed) * 100)
                        if (total_passed + total_failed) > 0 else 0, 1)

    summary_data = [
        ["Metric", "Value"],
        ["Total Runs", total_runs],
        ["Total Passed", total_passed],
        ["Total Failed", total_failed],
        ["Success Rate", f"{success_rate}%"],
        ["Exported At", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
    ]

    for row_idx, row in enumerate(summary_data, 1):
        for col_idx, value in enumerate(row, 1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=value)
            cell.border = border
            if row_idx == 1:
                cell.font = Font(bold=True, color="FFFFFF", size=12)
                cell.fill = PatternFill(start_color="7C3AED", end_color="7C3AED", fill_type="solid")

    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 30

    # Save
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def export_to_pdf(summaries):
    """Export reports to PDF"""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.enums import TA_CENTER

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=1.5 * cm,
        leftMargin=1.5 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CustomTitle",
        parent=styles["Heading1"],
        fontSize=20,
        textColor=colors.HexColor("#4F46E5"),
        alignment=TA_CENTER,
        spaceAfter=10,
    )

    elements = []

    # Title
    elements.append(Paragraph("ERP Test Automation - Reports", title_style))
    elements.append(Spacer(1, 0.3 * cm))

    # Summary
    total_runs = len(summaries)
    total_passed = sum(s.get("passed", 0) for s in summaries)
    total_failed = sum(s.get("failed", 0) + s.get("errors", 0) for s in summaries)
    success_rate = round((total_passed / (total_passed + total_failed) * 100)
                        if (total_passed + total_failed) > 0 else 0, 1)

    summary_text = f"""
    <b>Total Runs:</b> {total_runs} &nbsp;&nbsp; | &nbsp;&nbsp;
    <b>Passed:</b> {total_passed} &nbsp;&nbsp; | &nbsp;&nbsp;
    <b>Failed:</b> {total_failed} &nbsp;&nbsp; | &nbsp;&nbsp;
    <b>Success Rate:</b> {success_rate}%
    """
    elements.append(Paragraph(summary_text, styles["Normal"]))
    elements.append(Spacer(1, 0.5 * cm))

    # Table Data
    data = [["Run ID", "Test Name", "Total", "Passed", "Failed", "Duration (s)", "Timestamp"]]

    for s in summaries:
        failed = s.get("failed", 0) + s.get("errors", 0)
        data.append([
            s.get("run_id", "")[:20],
            s.get("test_name", "")[:25],
            str(s.get("total", 0)),
            str(s.get("passed", 0)),
            str(failed),
            str(int(s.get("duration_seconds", 0))),
            s.get("timestamp", "")[:15],
        ])

    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F46E5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 10),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 1), (-1, -1), 9),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
        ("TOPPADDING", (0, 0), (-1, 0), 8),
        ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#F8FAFC")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [colors.white, colors.HexColor("#F8FAFC")]),
    ]))

    elements.append(table)

    # Footer
    elements.append(Spacer(1, 0.5 * cm))
    footer = f"Exported at: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | ERP Test Automation v4.0"
    elements.append(Paragraph(footer, styles["Normal"]))

    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()
